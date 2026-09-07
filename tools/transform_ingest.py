from __future__ import annotations

import csv
import sys
import tempfile
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

# Turns an AWS Transform assessment export (the console ZIP: PPTX + XLSX + PDF, no JSON) into one
# text blob for `extract_contract`. The XLSX is the source of truth for numbers; the PDF/PPTX add
# the prose. All heavy parsers (openpyxl / pypdf / python-pptx) are imported lazily.

_MODELS = {"on_demand": "on-demand", "1yr_ri": "1yr", "3yr_ri": "3yr"}


@dataclass
class AssessmentPackage:
    servers: list[dict] = field(default_factory=list)
    volumes: list[dict] = field(default_factory=list)
    issues: list[dict] = field(default_factory=list)
    region: str = ""
    recommended_pricing: str = "on_demand"
    narrative: str = ""


def _num(v) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _pct(v: float | None) -> str:
    if v is None:
        return "n/a"
    return f"{v * 100:.1f}%" if v <= 1.5 else f"{v:.1f}%"


def _rows(ws):
    header = None
    for raw in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c).strip() if c is not None else "" for c in raw]
            continue
        yield dict(zip(header, raw))


def _as_dir(path: str | Path) -> Path:
    p = Path(path)
    if p.suffix.lower() != ".zip":
        return p
    dest = Path(tempfile.mkdtemp(prefix="ingest_"))
    with zipfile.ZipFile(p) as z:
        z.extractall(dest)
    return dest


def _read_csv(path: Path) -> list[dict]:
    if not path.is_file():
        return []
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def _pdf_text(pdf: Path) -> str:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError:
        return ""
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf)).pages)


def _pptx_text(pptx: Path) -> str:
    try:
        from pptx import Presentation
    except ModuleNotFoundError:
        return ""
    out = []
    for slide in Presentation(str(pptx)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
    return "\n".join(out)


def _recommended_pricing(text: str) -> str:
    t = text.lower()
    if "3-year reserved" in t or "3 year no upfront" in t:
        return "3yr_ri"
    if "1-year reserved" in t or "1 year no upfront" in t:
        return "1yr_ri"
    return "on_demand"


def _parse_workbook(xlsx: Path, pkg: AssessmentPackage) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    names = {s.title.lower(): s.title for s in wb.worksheets}

    by_id: dict[str, dict] = {}
    for model, needle in _MODELS.items():
        title = next(
            (t for lc, t in names.items() if lc.startswith("shared tenancy") and needle in lc), None
        )
        if title is None:
            continue
        for row in _rows(wb[title]):
            sid = row.get("Server Id")
            if not sid:
                continue
            rec = by_id.setdefault(
                sid,
                {
                    "server_id": sid,
                    "name": row.get("Server Name"),
                    "os": row.get("Operating System Name"),
                    "ram_gb": _num(row.get("RAM (GB)")),
                    "peak_cpu": _num(row.get("Peak CPU Utilization %")),
                    "avg_cpu": _num(row.get("Average CPU Utilization %")),
                    "peak_ram": _num(row.get("Peak RAM Utilization %")),
                    "ec2_recommended": row.get("EC2 Instance Recommended"),
                    "ec2_vcpu": _num(row.get("EC2 Total Cores")),
                    "ec2_ram_gb": _num(row.get("EC2 RAM (GB)")),
                    "region": row.get("AWS Region"),
                    "annualized_by_model": {},
                },
            )
            rec["annualized_by_model"][model] = _num(row.get("Annualized Total Cost"))
    pkg.servers = list(by_id.values())

    vtitle = names.get("block-ebs-cost-optimized")
    if vtitle:
        for row in _rows(wb[vtitle]):
            if not row.get("Volume Name"):
                continue
            pkg.volumes.append(
                {
                    "name": row.get("Volume Name"),
                    "service": row.get("AWS Service Mapping (Cost Optimized)"),
                    "provisioned_gib": _num(row.get("Capacity Total")),
                    "used_gib": _num(row.get("Capacity Used")),
                    "boot_gib": _num(row.get("VirtualBootVolumeSizeGB")),
                    "annual_usd": _num(row.get("Estimated Annual AWS Cost (Cost Optimized)")),
                    "region": row.get("AWS Region"),
                }
            )

    ftitle = names.get("file-fsxn")
    if ftitle:
        for row in _rows(wb[ftitle]):
            if not row.get("Volume Name"):
                continue
            pkg.volumes.append(
                {
                    "name": row.get("Volume Name"),
                    "service": "FSx for NetApp ONTAP",
                    "provisioned_gib": _num(row.get("Capacity Total (GiB)")),
                    "used_gib": _num(row.get("Capacity Used (GiB)")),
                    "boot_gib": None,
                    "annual_usd": None,  # FSx cost is only in the PDF Financial Summary
                    "region": None,
                }
            )

    ititle = names.get("assessment issues")
    if ititle:
        for row in _rows(wb[ititle]):
            if not row.get("Severity"):
                continue
            pkg.issues.append(
                {
                    "severity": row.get("Severity"),
                    "reason": row.get("Reason"),
                    "impact": row.get("Impact"),
                    "recommendation": row.get("Recommendation"),
                }
            )

    pkg.region = next(
        (s["region"] for s in pkg.servers if s.get("region")),
        next((v["region"] for v in pkg.volumes if v.get("region")), ""),
    )


def read_assessment(path: str | Path) -> AssessmentPackage:
    """`path` is a directory or the assessment .zip."""
    p = _as_dir(path)
    pkg = AssessmentPackage()
    pdf = next(iter(sorted(p.glob("*.pdf"))), None)
    pptx = next(iter(sorted(p.glob("*.pptx"))), None)
    xlsx = next(iter(sorted(p.glob("*.xlsx"))), None)

    if pdf is not None:
        pkg.narrative = _pdf_text(pdf)
    if not pkg.narrative and pptx is not None:
        pkg.narrative = _pptx_text(pptx)
    if xlsx is not None:
        _parse_workbook(xlsx, pkg)
    pkg.recommended_pricing = _recommended_pricing(pkg.narrative)
    return pkg


def _financial_summary(text: str) -> dict[str, dict[str, float]] | None:
    """Parse the PDF's Financial Summary rows (Compute / Storage / Network per pricing model).
    This is Transform's authoritative total - it includes FSx, which the XLSX volume sheets don't."""
    import re

    if "Financial Summary" not in text:
        return None
    out: dict[str, dict[str, float]] = {}
    for label, key in (("Compute", "compute"), ("Storage", "storage"), ("Network", "network")):
        m = re.search(
            rf"^\s*{label}\s+\$?([\d,]+)\s+\$?([\d,]+)\s+\$?([\d,]+)\s*$", text, re.MULTILINE
        )
        if not m:
            return None
        for model, v in zip(_MODELS, (float(g.replace(",", "")) for g in m.groups())):
            out.setdefault(model, {})[key] = v
    return out


def derived_cost_basis(pkg: AssessmentPackage) -> dict[str, dict[str, float]]:
    """Per pricing model: compute + network + storage, excluding AWS Business Support - the
    like-for-like number `monthly_run_rate` computes. Uses the PDF Financial Summary when present
    (authoritative, includes FSx), else sums the per-server / per-volume sheet costs."""
    fs = _financial_summary(pkg.narrative)
    out: dict[str, dict[str, float]] = {}
    for model in _MODELS:
        if fs and len(fs.get(model, {})) == 3:
            annual = round(sum(fs[model].values()), 2)
        else:
            compute = sum(s["annualized_by_model"].get(model) or 0.0 for s in pkg.servers)
            storage = sum(v.get("annual_usd") or 0.0 for v in pkg.volumes)
            annual = round(compute + storage, 2)
        out[model] = {"annual_usd": annual, "monthly_usd": round(annual / 12, 2)}
    return out


def to_business_case(pkg: AssessmentPackage) -> str:
    lines = ["# AWS Transform Migration Assessment", ""]
    if pkg.narrative:
        lines += ["## Assessment report (narrative)", "", pkg.narrative.strip(), ""]
    lines += [
        "## Structured facts (parsed from the analysis workbook)",
        "",
        f"Region: {pkg.region or 'unknown'}",
        f"Recommended pricing model: {pkg.recommended_pricing}",
        "",
        f"### Servers ({len(pkg.servers)})",
    ]
    for s in pkg.servers:
        costs = s["annualized_by_model"]
        cost_str = " / ".join(
            f"{m}=${costs[m]:,.2f}" for m in _MODELS if costs.get(m) is not None
        )
        lines.append(
            f"- {s['server_id']} ({s.get('name')}): {s.get('os')}, RAM {s.get('ram_gb')} GB, "
            f"peak CPU {_pct(s.get('peak_cpu'))}, avg CPU {_pct(s.get('avg_cpu'))}, "
            f"peak RAM {_pct(s.get('peak_ram'))}. Transform right-sizes to "
            f"{s.get('ec2_recommended')} ({s.get('ec2_vcpu')} vCPU, {s.get('ec2_ram_gb')} GiB). "
            f"Annualized total (compute+network): {cost_str}."
        )
    lines += ["", f"### Storage volumes ({len(pkg.volumes)})"]
    for v in pkg.volumes:
        cost = f" Annual cost ${v['annual_usd']:,.2f}." if v.get("annual_usd") is not None else ""
        boot = f", boot {v['boot_gib']} GiB" if v.get("boot_gib") is not None else ""
        lines.append(
            f"- {v['name']}: {v.get('service')}, {v.get('provisioned_gib')} GiB provisioned, "
            f"{v.get('used_gib')} GiB used{boot}.{cost}"
        )
    if pkg.issues:
        lines += ["", "### Assessment issues"]
        for i in pkg.issues:
            lines.append(
                f"- [{i.get('severity')}] {i.get('reason')} -> {i.get('impact')}. "
                f"{i.get('recommendation')}"
            )
    lines += [
        "",
        "### Derived cost basis (compute + network + storage; excludes AWS Business Support)",
    ]
    for model, c in derived_cost_basis(pkg).items():
        lines.append(
            f"- {model}: ${c['annual_usd']:,.2f}/yr  (${c['monthly_usd']:,.2f}/month)"
        )
    return "\n".join(lines) + "\n"


# --- Discovery-tool export (the CSVs you upload to Transform / MPA). Per the plan the discovery
# CSV is another team's, but a business case is "free-form", so the ingest handles it too. ---

_APP_KEYWORDS = (
    "java", "python", "node", "redis", "sqlservr", "mssql", "oracle", "postgres", "mysql",
    "mariadb", "mongo", "nginx", "apache2", "httpd", "haproxy", "dockerd", "containerd",
    "jenkins", "nfsd", "rpc.idmapd", "rabbitmq", "kafka", "activemq", "tomcat", "dotnet",
    "gunicorn", "uvicorn", "php", "RSHostingService", "launchpad",
)


@dataclass
class DiscoveryPackage:
    servers: list[dict] = field(default_factory=list)
    dependencies: list[dict] = field(default_factory=list)
    databases: list[dict] = field(default_factory=list)


def _looks_like_app(name: str, cmd: str) -> bool:
    blob = f"{name} {cmd}".lower()
    if "===section:" in blob or "estate-chatter" in blob or "ssm-agent" in blob:
        return False
    return any(k.lower() in blob for k in _APP_KEYWORDS)


def read_discovery(path: str | Path) -> DiscoveryPackage:
    root = _as_dir(path)
    base = root / "full_exports" if (root / "full_exports").is_dir() else root
    if not (base / "server_inventory.csv").is_file():
        found = next(iter(root.glob("**/server_inventory.csv")), None)
        base = found.parent if found else base

    inv = _read_csv(base / "server_inventory.csv")
    perf = {r["server_id"]: r for r in _read_csv(base / "server_performance_metrics.csv")}
    name_by_id = {r["server_id"]: r.get("server_name") or r["server_id"] for r in inv}

    disk: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    for r in _read_csv(base / "server_storage_performance.csv"):
        d = disk[r["server_id"]]
        d[0] += _num(r.get("disk_total_gb")) or 0.0
        d[1] += _num(r.get("disk_used_gb")) or 0.0

    apps: dict[str, set[str]] = defaultdict(set)
    for r in _read_csv(base / "process_metrics.csv"):
        nm, cmd = (r.get("process_name") or "").strip(), (r.get("process_command_line") or "").strip()
        if nm and not nm.startswith("[") and _looks_like_app(nm, cmd):
            apps[r["server_id"]].add(cmd or nm)

    servers = []
    for r in inv:
        sid = r["server_id"]
        pr = perf.get(sid, {})
        servers.append({
            "server_id": sid,
            "name": r.get("server_name"),
            "os": f"{r.get('os_name', '')} {r.get('os_version', '')}".strip(),
            "os_type": r.get("os_type"),
            "cpu": _num(r.get("cpu_count")),
            "ram_gb": _num(r.get("total_memory_gb")),
            "disk_total_gb": round(disk[sid][0], 1),
            "disk_used_gb": round(disk[sid][1], 1),
            "ip": r.get("primary_ip_address"),
            "environment": r.get("environment") or "",
            "cpu_p95": _num(pr.get("cpu_utilization_p95_pct")),
            "cpu_peak": _num(pr.get("cpu_utilization_peak_pct")),
            "mem_p95": _num(pr.get("memory_utilization_p95_pct")),
            "mem_peak": _num(pr.get("memory_utilization_peak_pct")),
            "apps": sorted(apps.get(sid, []))[:8],
        })

    deps = []
    for r in _read_csv(base / "network_data_full.csv"):
        deps.append({
            "src": name_by_id.get(r.get("Source Server ID"), r.get("Source Server ID")),
            "dst": name_by_id.get(r.get("Target Server ID"), r.get("Target Server ID")),
            "port": r.get("Target Port"),
            "proto": r.get("Transport Protocol"),
            "src_process": r.get("Source Process Name"),
            "count": _num(r.get("Count")),
        })

    dbs = []
    for r in _read_csv(base / "sql_server_data_full.csv"):
        if (r.get("Is Engine Component") or "").strip().upper() != "Y":
            continue
        dbs.append({
            "server": name_by_id.get(r.get("Server ID"), r.get("Server ID")),
            "engine": "SQL Server",
            "version": r.get("Version"),
            "edition": r.get("Edition"),
            "instance": r.get("Instance Name"),
        })
    for r in _read_csv(base / "oracle_data_cdbs_full.csv"):
        dbs.append({
            "server": name_by_id.get(r.get("Server ID"), r.get("Server ID")),
            "engine": "Oracle",
            "version": r.get("Version Full") or r.get("Version"),
            "edition": r.get("Edition"),
            "instance": r.get("DB Name"),
        })
    return DiscoveryPackage(servers, deps, dbs)


def _discovery_notes(pkg: DiscoveryPackage) -> list[str]:
    notes = []
    for db in pkg.databases:
        ed = (db.get("edition") or "").lower()
        if "express" in ed or db.get("edition") == "XE":
            notes.append(
                f"{db['engine']} on {db['server']} is a license-limited edition "
                f"({db.get('edition')}) - capacity/feature ceilings apply on migration."
            )
    for s in pkg.servers:
        p95, peak = s.get("cpu_p95") or 0, s.get("cpu_peak") or 0
        if peak >= 99 and p95 >= 95:
            tail = " and no application process was captured" if not s.get("apps") else ""
            notes.append(f"{s['name']} hits 100% CPU (p95 {p95:.0f}%){tail} - review sizing.")
        if (s.get("mem_p95") or 0) >= 85:
            notes.append(f"{s['name']} is memory-constrained (mem p95 {s['mem_p95']:.0f}%).")
        if s.get("os_type") == "Windows":
            notes.append(f"{s['name']} is a Windows workload ({s.get('os')}).")
        if any("docker" in a or "containerd" in a for a in s.get("apps", [])):
            notes.append(f"{s['name']} already runs containers on the source host.")
    return notes


def discovery_to_business_case(pkg: DiscoveryPackage) -> str:
    lines = [
        "# Discovery Tool Export",
        "",
        "## Structured facts (parsed from the discovery export)",
        "",
        f"### Servers ({len(pkg.servers)})",
    ]
    for s in pkg.servers:
        apps = "; ".join(s["apps"]) or "(no notable application processes captured)"
        lines.append(
            f"- {s['name']} ({s.get('ip')}): {s.get('os')}, {s.get('cpu')} vCPU, "
            f"{s.get('ram_gb')} GB RAM, disk {s.get('disk_used_gb')}/{s.get('disk_total_gb')} GB. "
            f"CPU p95 {_pct(s.get('cpu_p95'))} peak {_pct(s.get('cpu_peak'))}, "
            f"mem p95 {_pct(s.get('mem_p95'))} peak {_pct(s.get('mem_peak'))}. "
            f"env={s.get('environment') or 'n/a'}. processes: {apps}"
        )
    lines += ["", f"### Databases ({len(pkg.databases)})"]
    for d in pkg.databases:
        lines.append(
            f"- {d['server']}: {d['engine']} {d.get('version')} {d.get('edition')}, "
            f"instance {d.get('instance')}"
        )
    lines += ["", f"### Dependencies ({len(pkg.dependencies)} observed connections)"]
    for d in pkg.dependencies:
        lines.append(
            f"- {d['src']} --[{d.get('proto')} {d.get('port')}, {d.get('src_process') or '?'}, "
            f"{int(d.get('count') or 0)} conns]--> {d['dst']}"
        )
    notes = _discovery_notes(pkg)
    if notes:
        lines += ["", "### Notes (auto-derived)"]
        lines += [f"- {n}" for n in dict.fromkeys(notes)]
    return "\n".join(lines) + "\n"


# --- Migration plan (the wave-planning export: MGN import + enriched inventory + apps) ---

# vCPU / GiB for the instance types the wave plan assigns, so the plan's own sizing can be checked
# against the inventory it was built from.
INSTANCE_SPEC = {
    "t3a.nano": (2, 0.5),
    "t3a.micro": (2, 1.0),
    "t3a.small": (2, 2.0),
    "t2.small": (1, 2.0),
    "t3.micro": (2, 1.0),
    "t3.small": (2, 2.0),
    "c5a.large": (2, 4.0),
    "m7a.medium": (1, 4.0),
    "m5.large": (2, 8.0),
}


@dataclass
class MigrationPlanPackage:
    wave: str = ""
    servers: list[dict] = field(default_factory=list)
    apps: list[dict] = field(default_factory=list)
    databases: list[dict] = field(default_factory=list)
    move_groups: dict[str, str] = field(default_factory=dict)
    citations: dict[str, dict] = field(default_factory=dict)
    audit: dict = field(default_factory=dict)


def _one(root: Path, *patterns: str) -> Path | None:
    for pat in patterns:
        hits = sorted(root.glob(pat)) or sorted(root.glob(f"**/{pat}"))
        if hits:
            return hits[-1]
    return None


def read_migration_plan(path: str | Path) -> MigrationPlanPackage:
    root = _as_dir(path)
    pkg = MigrationPlanPackage()

    by_id: dict[str, dict] = {}
    for r in _read_csv(_one(root, "mgn_import*.csv") or Path()):
        sid = r.get("mgn:server:user-provided-id") or r.get("mgn:server:tag:hostname") or ""
        if not sid:
            continue
        pkg.wave = pkg.wave or (r.get("mgn:wave:name") or "")
        by_id[sid] = {
            "name": sid,
            "ip": r.get("mgn:launch:nic:0:private-ip:0"),
            "platform": r.get("mgn:server:platform"),
            "instance_type": r.get("mgn:launch:instance-type"),
            "tenancy": r.get("mgn:launch:placement:tenancy"),
            "app_id": r.get("mgn:app:name"),
            "vmmoref": r.get("mgn:server:tag:vmmoref"),
        }

    for r in _read_csv(_one(root, "enriched_inventory*.csv") or Path()):
        sid = r.get("mgn:server:user-provided-id") or ""
        s = by_id.setdefault(sid, {"name": sid})
        mib = _num(r.get("metadata:total_system_memory_mebibytes"))
        s.update({
            "app_name": r.get("metadata:app_name"),
            "os": r.get("metadata:operating_system_type"),
            "cores": _num(r.get("metadata:total_cpu_cores")),
            "ram_gib": round(mib / 1024, 2) if mib else None,
            "cpu_avg": _num(r.get("metadata:average_cpu_utilization_percentage")),
            "cpu_peak": _num(r.get("metadata:peak_cpu_utilization_percentage")),
            "mem_avg": _num(r.get("metadata:average_memory_utilization_percentage")),
            "mem_peak": _num(r.get("metadata:peak_memory_utilization_percentage")),
            "move_group_id": r.get("metadata:move_group_id"),
            "move_group": r.get("metadata:move_group_name"),
            "reason": r.get("metadata:reason"),
        })
    pkg.servers = sorted(by_id.values(), key=lambda s: s["name"])

    for r in _read_csv(_one(root, "applications*.csv") or Path()):
        if "server_ids" not in r:
            continue
        pkg.apps.append({
            "id": r.get("app_id"),
            "name": r.get("app_name"),
            "strategy": r.get("migration_strategy"),
            "environment": r.get("environment"),
            "confidence": _num(r.get("confidence_score")),
            "server_count": int(_num(r.get("server_count")) or 0),
            "priority_rank": r.get("priority_rank"),
            "scoring": r.get("scoring_justification"),
        })
    for r in _read_csv(_one(root, "databases*.csv") or Path()):
        if "server_ids" in r:
            continue
        pkg.databases.append({
            "id": r.get("app_id"), "name": r.get("app_name"),
            "strategy": r.get("migration_strategy"),
        })

    for r in _read_csv(_one(root, "apps_to_move_groups*.csv") or Path()):
        if r.get("app_id"):
            pkg.move_groups[r["app_id"]] = r.get("move_group_name") or ""

    for r in _read_csv(_one(root, "citations*.csv") or Path()):
        eid = r.get("entity_id")
        if eid:
            pkg.citations[eid] = {
                "source": r.get("source"), "row": r.get("source_row"),
                "move_group_rule": r.get("move_group_rule"), "wave_rule": r.get("wave_rule"),
            }

    audit = _one(root, "audit*.json")
    if audit and audit.is_file():
        import json

        pkg.audit = json.loads(audit.read_text(encoding="utf-8"))
    return pkg


def plan_sizing_findings(pkg: MigrationPlanPackage) -> list[dict]:
    """Where the wave plan's own instance choice is smaller than the inventory it was built from."""
    out = []
    for s in pkg.servers:
        spec = INSTANCE_SPEC.get(s.get("instance_type") or "")
        if not spec:
            continue
        vcpu, gib = spec
        short = []
        if s.get("cores") and vcpu < s["cores"]:
            short.append(f"{vcpu} vCPU < {s['cores']:.0f} cores observed")
        if s.get("ram_gib") and gib < s["ram_gib"]:
            short.append(f"{gib} GiB RAM < {s['ram_gib']} GiB observed")
        if short:
            out.append({
                "server": s["name"], "instance_type": s["instance_type"],
                "shortfalls": short, "cpu_peak": s.get("cpu_peak"), "mem_peak": s.get("mem_peak"),
            })
    return out


def plan_run_rate(pkg: MigrationPlanPackage, pricing_model: str = "on_demand") -> float:
    """The wave's monthly cost as the plan sizes it. Arithmetic, so the model never guesses it."""
    from tools.finops import monthly_run_rate

    resources = [{"instance_type": s.get("instance_type") or ""} for s in pkg.servers]
    return monthly_run_rate(resources, pricing_model=pricing_model)


def _plan_config_line(cfg: dict) -> str:
    cap = (cfg.get("size_constraints") or {}).get("max_resources_per_entity")
    return (
        f"- grouped by {cfg.get('group_by_attributes')}, priority {cfg.get('priority')}, "
        f"max {cap} servers per wave, seed {cfg.get('seed')}"
    )


def migration_plan_to_business_case(pkg: MigrationPlanPackage) -> str:
    lines = [
        f"# AWS Transform Migration Plan - {pkg.wave or 'wave'}",
        "",
        "## Structured facts (parsed from the wave-planning export)",
        "",
        f"### Applications ({len(pkg.apps)})",
    ]
    for a in pkg.apps:
        mg = pkg.move_groups.get(a["id"] or "", "n/a")
        lines.append(
            f"- {a['name']} [{a['id']}]: strategy={a['strategy']}, env={a['environment']}, "
            f"{a['server_count']} servers, move group {mg}, confidence {a['confidence']}, "
            f"priority rank {a['priority_rank']}"
        )
    if pkg.databases:
        lines += ["", f"### Database components ({len(pkg.databases)})"]
        lines += [f"- {d['name']} [{d['id']}]: strategy={d['strategy']}" for d in pkg.databases]

    lines += ["", f"### Servers in {pkg.wave or 'the wave'} ({len(pkg.servers)})"]
    for s in pkg.servers:
        lines.append(
            f"- {s['name']} ({s.get('ip')}): {s.get('os')}, planned instance "
            f"{s.get('instance_type')}, observed {s.get('cores'):.0f} cores / "
            f"{s.get('ram_gib')} GiB. CPU avg {_pct(s.get('cpu_avg'))} peak {_pct(s.get('cpu_peak'))}, "
            f"mem avg {_pct(s.get('mem_avg'))} peak {_pct(s.get('mem_peak'))}. "
            f"app={s.get('app_name')}, move group={s.get('move_group')}"
            if s.get("cores") else f"- {s['name']}: planned instance {s.get('instance_type')}"
        )

    findings = plan_sizing_findings(pkg)
    if findings:
        head = f"### Plan sizing shortfalls (auto-derived, {len(findings)} of {len(pkg.servers)} servers)"
        lines += ["", head]
        for f_ in findings:
            lines.append(
                f"- {f_['server']}: plan assigns {f_['instance_type']} but "
                + "; ".join(f_["shortfalls"])
                + f" (peak CPU {_pct(f_['cpu_peak'])}, peak mem {_pct(f_['mem_peak'])})"
            )

    rate = plan_run_rate(pkg)
    if rate:
        line = f"- on-demand run-rate of the {len(pkg.servers)} planned instances: ${rate:,.2f}/month"
        lines += ["", "### Derived cost basis (arithmetic, from the planned instance types)", line]

    v = (pkg.audit.get("validator_results") or {}) if pkg.audit else {}
    if v:
        lines += ["", "### Plan validators"]
        lines += [f"- {k}: {val}" for k, val in v.items()]
    cfg = (pkg.audit.get("config") or {}) if pkg.audit else {}
    if cfg:
        lines += [
            "",
            "### Plan configuration",
            _plan_config_line(cfg),
        ]
    if pkg.citations:
        src = {c["source"] for c in pkg.citations.values() if c.get("source")}
        lines += ["", "### Provenance",
                  f"- {len(pkg.citations)} entities cited from: {', '.join(sorted(src))}"]
    return "\n".join(lines) + "\n"


# The wave-plan artifacts the parser above needs, by fileMetadata.path prefix.
PLAN_ARTIFACTS = (
    "mgn_import", "enriched_inventory", "applications", "apps_to_move_groups", "citations", "audit"
)
ARTIFACT_CACHE = ".transform_cache"


def read_migration_plan_from_mcp(workspace, job_id: str, dest: str | Path | None = None):
    """Pull the wave plan straight out of the Transform workspace. Nothing about the estate is
    hard-coded: the artifacts, their ids and their contents all come from the service."""
    root = Path(dest or Path(ARTIFACT_CACHE) / job_id)
    workspace.fetch_all(job_id, root, names=PLAN_ARTIFACTS)
    return read_migration_plan(root)


def business_case_from_mcp(workspace, job_id: str, dest: str | Path | None = None) -> str:
    return migration_plan_to_business_case(read_migration_plan_from_mcp(workspace, job_id, dest))


def load_business_case(path: str | Path) -> str:
    """Auto-detect a migration plan (MGN import CSV), a discovery export (server_inventory.csv) or a
    Transform assessment (PPTX/XLSX/PDF), and flatten any of them into the prose blob
    `extract_contract` takes."""
    root = _as_dir(path)
    if _one(root, "mgn_import*.csv"):
        return migration_plan_to_business_case(read_migration_plan(root))
    if (root / "server_inventory.csv").is_file() or list(root.glob("**/server_inventory.csv")):
        return discovery_to_business_case(read_discovery(root))
    return to_business_case(read_assessment(root))


def main(argv: list[str] | None = None) -> int:
    argv = argv or sys.argv[1:]
    if not argv:
        print("usage: python -m tools.transform_ingest <plan | assessment.zip | discovery.zip | dir>")
        return 1
    print(load_business_case(argv[0]))
    return 0


if __name__ == "__main__":
    sys.exit(main())


def wave_sizing(pkg: MigrationPlanPackage) -> list[dict]:
    """What the plan would actually launch, per server, so the policy engine can rule on it."""
    out = []
    for s in pkg.servers:
        vcpu, gib = INSTANCE_SPEC.get(s.get("instance_type") or "", (None, None))
        out.append({
            "name": s.get("name"), "instance_type": s.get("instance_type"),
            "vcpu": vcpu, "ram_gib": gib,
        })
    return out


def wave_step_params(pkg: MigrationPlanPackage) -> dict:
    """What the deterministic steps need, taken from the plan instead of typed in by hand."""
    return {
        "source_server_ids": [s["name"] for s in pkg.servers if s.get("name")],
        "wave": pkg.wave,
    }
