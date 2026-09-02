from pathlib import Path

import openpyxl

from tools.transform_ingest import (
    derived_cost_basis,
    discovery_to_business_case,
    load_business_case,
    read_assessment,
    read_discovery,
    to_business_case,
)

_COST = {"On-Demand": 479.3, "1yr NU": 327.14, "3yr NU": 227.63}


def _make_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for model in ("On-Demand", "1yr NU", "3yr NU"):
        ws = wb.create_sheet(f"Shared Tenancy - {model}")
        ws.append([
            "Server Id", "Server Name", "RAM (GB)", "Operating System Name",
            "Peak CPU Utilization %", "Average CPU Utilization %", "Peak RAM Utilization %",
            "EC2 Instance Recommended", "EC2 Total Cores", "EC2 RAM (GB)", "AWS Region",
            "Annualized Total Cost",
        ])
        ws.append([
            "s1", "host1", 2, "Ubuntu 16.04.7 LTS", 0.48, 0.17, 0.17,
            "c7a.medium", 1, 2, "us-east-1", _COST[model],
        ])
    v = wb.create_sheet("Block-EBS-Cost-Optimized")
    v.append([
        "Volume Name", "AWS Service Mapping (Cost Optimized)", "AWS Region", "Capacity Total",
        "Capacity Used", "VirtualBootVolumeSizeGB", "Estimated Annual AWS Cost (Cost Optimized)",
    ])
    v.append(["s1-storage", "EBS GP3", "us-east-1", 30, 2.7, 30, 29.76])
    f = wb.create_sheet("File-FSxN")
    f.append(["Volume Name", "Capacity Total (GiB)", "Capacity Used (GiB)"])
    f.append(["s1-nas", 29.93, 1.95])
    i = wb.create_sheet("Assessment Issues")
    i.append(["Resource ID", "Resource Name", "Resource Type", "Severity", "Reason", "Impact", "Recommendation"])
    i.append(["-", "-", "-", "WARNING", "missing host data", "sustainability excluded", "provide host data"])
    wb.save(path)


def test_reads_servers_volumes_issues(tmp_path):
    _make_xlsx(tmp_path / "analysis.xlsx")
    pkg = read_assessment(tmp_path)
    assert len(pkg.servers) == 1
    s = pkg.servers[0]
    assert s["os"] == "Ubuntu 16.04.7 LTS" and s["ec2_recommended"] == "c7a.medium"
    assert s["annualized_by_model"] == {"on_demand": 479.3, "1yr_ri": 327.14, "3yr_ri": 227.63}
    ebs = next(v for v in pkg.volumes if v["service"] == "EBS GP3")
    fsx = next(v for v in pkg.volumes if v["service"] == "FSx for NetApp ONTAP")
    assert ebs["annual_usd"] == 29.76 and fsx["annual_usd"] is None and fsx["provisioned_gib"] == 29.93
    assert pkg.issues[0]["severity"] == "WARNING"
    assert pkg.region == "us-east-1"


def test_derived_cost_basis_sums_sheet_costs_without_a_financial_summary(tmp_path):
    _make_xlsx(tmp_path / "a.xlsx")  # synthetic workbook, no PDF -> no Financial Summary
    basis = derived_cost_basis(read_assessment(tmp_path))
    assert basis["3yr_ri"]["annual_usd"] == round(227.63 + 29.76, 2)  # FSx has no cost here


def test_financial_summary_parser_and_precedence():
    from tools.transform_ingest import AssessmentPackage, _financial_summary

    narrative = (
        "Financial Summary\n"
        "Cost Component AWS - On Demand AWS - 1 Year No Upfront AWS - 3 Year No Upfront\n"
        "Compute $9,654 $6,763 $5,076\n"
        "Storage $1,447 $1,447 $1,447\n"
        "Network $515 $515 $515\n"
        "Business Support $1,162 $872 $704\n"
    )
    fs = _financial_summary(narrative)
    assert fs["3yr_ri"] == {"compute": 5076.0, "storage": 1447.0, "network": 515.0}
    pkg = AssessmentPackage(narrative=narrative)
    # Financial Summary wins over (here empty) sheet sums; Business Support is excluded
    assert derived_cost_basis(pkg)["3yr_ri"] == {"annual_usd": 7038.0, "monthly_usd": 586.5}


def test_business_case_text_carries_the_facts(tmp_path):
    _make_xlsx(tmp_path / "a.xlsx")
    bc = to_business_case(read_assessment(tmp_path))
    for token in ("Ubuntu 16.04.7 LTS", "c7a.medium", "us-east-1", "Derived cost basis", "WARNING"):
        assert token in bc


def test_the_committed_real_fixture_matches():
    bc = (Path(__file__).resolve().parents[1] / "fixtures" / "vmware-001" / "business_case.md").read_text(
        encoding="utf-8"
    )
    assert "server-f5a2cb54b18d7337" in bc
    # from the PDF Financial Summary (authoritative), not the per-line sheet sums
    assert "$515.00/yr" in bc and "$42.92/month" in bc


def test_financial_summary_and_fsx_from_the_committed_mixed_estate_fixture():
    bc = (
        Path(__file__).resolve().parents[1] / "fixtures" / "mixed-estate-001" / "business_case.md"
    ).read_text(encoding="utf-8")
    assert "FSx for NetApp ONTAP" in bc  # the NFS box, from the File-FSxN sheet
    assert "3yr_ri: $7,038.00/yr  ($586.50/month)" in bc  # compute+network+storage incl. FSx
    assert "vCPU specification missing or invalid" in bc  # the 12 provisional-sizing warnings


# --- discovery-tool export ---


def _make_discovery(d: Path) -> None:
    (d / "server_inventory.csv").write_text(
        "server_id,server_name,os_type,os_name,os_version,cpu_count,total_memory_gb,"
        "primary_ip_address,environment\n"
        "s1,catalog-svc-01,Linux,Amazon Linux,2023,2,1.86,10.0.0.1,prod\n"
        "s2,db-01,Linux,Amazon Linux,2023,2,3.75,10.0.0.2,prod\n",
        encoding="utf-8",
    )
    (d / "server_performance_metrics.csv").write_text(
        "server_id,cpu_utilization_p95_pct,cpu_utilization_peak_pct,memory_utilization_p95_pct,"
        "memory_utilization_peak_pct\ns1,95.0,100.0,18.0,19.0\ns2,100.0,100.0,66.0,67.0\n",
        encoding="utf-8",
    )
    (d / "server_storage_performance.csv").write_text(
        "server_id,disk_total_gb,disk_used_gb\ns1,30,2.2\ns2,60,6.8\n", encoding="utf-8"
    )
    (d / "process_metrics.csv").write_text(
        "server_id,process_name,process_command_line\n"
        "s1,java,java -jar /opt/catalog/catalog.jar\n"
        "s1,systemd,/usr/lib/systemd/systemd\n"
        "s2,tnslsnr,/opt/oracle/product/21c/dbhomeXE/bin/tnslsnr LISTENER\n",
        encoding="utf-8",
    )
    (d / "network_data_full.csv").write_text(
        "Source Server ID,Target Server ID,Target Port,Transport Protocol,Source Process Name,Count\n"
        "s1,s2,1521,TCP,java,2993\n",
        encoding="utf-8",
    )
    (d / "oracle_data_cdbs_full.csv").write_text(
        "Server ID,DB Name,Version,Version Full,Edition\ns2,XE,21.0.0.0.0,21.3.0.0.0,XE\n",
        encoding="utf-8",
    )


def test_read_discovery(tmp_path):
    _make_discovery(tmp_path)
    pkg = read_discovery(tmp_path)
    assert len(pkg.servers) == 2
    s1 = next(s for s in pkg.servers if s["name"] == "catalog-svc-01")
    assert "java -jar /opt/catalog/catalog.jar" in s1["apps"]
    assert "systemd" not in " ".join(s1["apps"])
    assert s1["disk_used_gb"] == 2.2 and s1["disk_total_gb"] == 30.0
    assert len(pkg.dependencies) == 1 and pkg.dependencies[0]["port"] == "1521"
    assert pkg.databases[0]["engine"] == "Oracle" and pkg.databases[0]["edition"] == "XE"


def test_discovery_business_case_and_notes(tmp_path):
    _make_discovery(tmp_path)
    bc = discovery_to_business_case(read_discovery(tmp_path))
    assert "catalog-svc-01" in bc and "Oracle 21.3.0.0.0 XE" in bc
    assert "--[TCP 1521, java, 2993 conns]-->" in bc
    assert "license-limited edition (XE)" in bc


def test_load_business_case_detects_a_discovery_export(tmp_path):
    _make_discovery(tmp_path)
    assert load_business_case(tmp_path).startswith("# Discovery Tool Export")


def test_committed_discovery_fixture_matches():
    bc = (
        Path(__file__).resolve().parents[1] / "fixtures" / "discovery-001" / "business_case.md"
    ).read_text(encoding="utf-8")
    assert "Oracle 21.3.0.0.0 XE" in bc and "SQL Server 2022.160.4265.3 Express" in bc
    assert "EC2AMAZ-HVB61P6 hits 100% CPU" in bc
